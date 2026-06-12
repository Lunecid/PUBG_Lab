from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/lunecid/project/PUBG/network_analysis/PUBG_safe_zone_STDB_research_proposal.docx"

FONT = "Malgun Gothic"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(89, 89, 89)
LIGHT_FILL = "F4F6F9"
BORDER = "B8C2CC"


def set_run_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Pt(widths[i] / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[i]))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def paragraph_border_bottom(paragraph, color="A6A6A6", size="6", space="4"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def add_para(doc, text="", style=None, align=None, before=None, after=None, line_spacing=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run_font(r)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, size=16, bold=True, color=BLUE)
    elif level == 2:
        set_run_font(r, size=13, bold=True, color=BLUE)
    else:
        set_run_font(r, size=12, bold=True, color=DARK_BLUE)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_shading(hdr[i], LIGHT_FILL)
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, bold=True, size=9.5)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=9.5)
            if len(value) < 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def apply_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.28

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("PUBG Safe Zone 기반 시공간 생존환경 분석 연구계획서")
    set_run_font(run, size=9, color=GRAY)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("PUBG Telemetry의 시공간 데이터베이스 모델링과\nSafe Zone 기반 생존환경 분석")
    set_run_font(r, size=21, bold=True, color=RGBColor(0, 0, 0))

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(22)
    r = p2.add_run("Research Proposal")
    set_run_font(r, size=12, color=GRAY)

    meta = [
        ("문서 유형", "연구계획서"),
        ("연구 대상", "PUBG 공식 텔레메트리"),
        ("연구 방법", "시공간 데이터베이스 모델링 및 질의 기반 분석"),
        ("작성일", "2026년 5월 24일"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    set_table_width(table, [1900, 7460])
    set_table_borders(table, color="D9DEE5", size="4")
    for i, (label, value) in enumerate(meta):
        cells = table.rows[i].cells
        set_cell_shading(cells[0], LIGHT_FILL)
        p = cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, size=9.5, bold=True)
        p = cells[1].paragraphs[0]
        r = p.add_run(value)
        set_run_font(r, size=9.5)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(16)
    rule.paragraph_format.space_after = Pt(12)
    paragraph_border_bottom(rule, color="A6A6A6", size="8")


def main():
    doc = Document()
    apply_styles(doc)
    add_title_block(doc)

    add_heading(doc, "초록", 1)
    add_para(
        doc,
        "본 연구는 PUBG 경기 텔레메트리를 시공간 데이터베이스 관점에서 모델링하고, "
        "안전구역 축소에 따른 팀의 이동궤적과 생존환경 변화를 분석하는 것을 목적으로 한다. "
        "PUBG 경기에서는 시간이 지남에 따라 안전구역이 줄어들고, 생존한 팀들은 점차 좁은 공간으로 이동한다. "
        "이 과정에서 팀 간 거리, 주변 적 밀도, 교전 노출 정도가 함께 변하며, 이러한 변화는 최종 등수와 관련될 수 있다. "
        "본 연구는 플레이어와 팀의 위치를 이동 객체로, 안전구역을 시간에 따라 변하는 공간 영역으로, "
        "전투 이벤트를 특정 시간과 위치에서 발생하는 공간 사건으로 정리한다. 이후 안전구역 거리 질의, 반경 질의, "
        "최근접 질의, 이동궤적 질의, 시간창 질의를 이용하여 팀별 생존환경 지표를 생성하고, 이 지표들이 최종 등수와 "
        "어떤 관계를 갖는지 분석한다.",
    )
    p = add_para(doc, "", after=12)
    r = p.add_run("주요어: ")
    set_run_font(r, bold=True)
    r = p.add_run("PUBG, 텔레메트리, 시공간 데이터베이스, 안전구역, 이동궤적, 생존환경")
    set_run_font(r)

    add_heading(doc, "1. 서론", 1)
    add_para(
        doc,
        "온라인 게임 로그는 단순한 행동 기록이 아니라 시간과 공간이 함께 포함된 데이터이다. "
        "특히 PUBG와 같은 배틀로얄 게임에서는 플레이어가 지도 위를 이동하고, 시간이 지나면서 안전구역이 축소되며, "
        "전투 이벤트가 특정 위치에서 발생한다. 이러한 특징 때문에 PUBG 텔레메트리는 시공간 데이터베이스의 관점에서 "
        "분석하기에 적합하다.",
    )
    add_para(
        doc,
        "기존의 게임 분석은 주로 경기 종료 후 집계되는 통계를 사용한다. 킬 수, 총 피해량, 이동 거리, 생존 시간은 "
        "팀의 활동량을 요약할 수 있지만, 특정 시점에 팀이 어떤 위치에 있었는지, 안전구역 경계와 얼마나 가까웠는지, "
        "주변에 적 팀이 얼마나 있었는지를 충분히 보여주지는 못한다. 따라서 경기 결과를 더 구체적으로 설명하기 위해서는 "
        "위치, 시간, 사건을 함께 다룰 수 있는 분석 구조가 필요하다.",
    )
    add_para(
        doc,
        "본 연구는 PUBG 텔레메트리를 시공간 데이터베이스로 구조화하고, 안전구역 축소에 따라 팀이 놓이는 생존환경이 "
        "어떻게 달라지는지 분석한다. 분석의 초점은 전략 유형 분류나 복잡한 예측 모델이 아니라, 안전구역을 기준으로 "
        "팀의 위치와 주변 환경을 설명할 수 있는 지표를 정의하고, 이 지표가 최종 등수와 어떤 관계를 갖는지 확인하는 데 있다.",
    )

    add_heading(doc, "2. 선행연구 및 연구의 차별성", 1)
    add_para(
        doc,
        "Sagong, Woo, Kim의 연구는 PUBG에서 단계별 전략과 생존 예측을 다루었다. 이 연구는 플레이어의 행동을 "
        "공격성과 활동성으로 나누고, 이를 바탕으로 Hunter, Sniper, Traveler, Camper와 같은 전략 유형을 정의하였다. "
        "이후 각 전략 유형이 다음 단계의 생존 여부와 최종 승리 예측에 어떤 영향을 주는지 분석하였다.",
    )
    add_para(
        doc,
        "이 선행연구는 PUBG에서 단계별 분석이 중요하다는 점을 보여준다. 또한 안전구역이 축소되면서 플레이어의 행동이 "
        "달라질 수 있고, 단계별 특성이 생존 결과와 연결될 수 있음을 제시한다. 본 연구도 안전구역 변화와 단계별 분석을 "
        "중요하게 본다는 점에서 이 연구와 관련된다.",
    )
    add_para(
        doc,
        "그러나 본 연구의 초점은 전략 유형 분류가 아니다. 본 연구는 PUBG 텔레메트리를 시공간 데이터베이스로 구조화하고, "
        "안전구역을 기준으로 팀이 어떤 위치 환경에 놓였는지를 분석한다. 기존 연구가 플레이어의 행동 유형을 중심으로 "
        "생존을 설명했다면, 본 연구는 축소되는 안전구역 안에서 팀이 처한 시공간적 생존환경을 중심으로 최종 등수를 분석한다.",
    )

    add_heading(doc, "3. 연구 질문", 1)
    add_para(
        doc,
        "본 연구는 세 가지 질문을 중심으로 진행한다. 첫째, PUBG 텔레메트리는 어떤 시공간 데이터베이스 구조로 표현될 수 "
        "있는가. 둘째, 안전구역 축소에 따라 팀의 이동궤적과 생존환경은 어떻게 변화하는가. 셋째, 시공간 분석을 통해 얻은 "
        "생존환경 요소 중 최종 등수와 관련이 큰 것은 무엇인가.",
    )
    add_para(
        doc,
        "첫 번째 질문은 데이터 모델링에 관한 것이다. 플레이어와 팀의 위치는 이동 객체로 볼 수 있고, 안전구역은 시간에 따라 "
        "변하는 공간 영역으로 볼 수 있으며, 전투 이벤트는 특정 시간과 위치에서 발생한 사건으로 볼 수 있다. 두 번째 질문은 "
        "안전구역 변화에 대한 팀의 반응을 다룬다. 세 번째 질문은 최종 등수와 관련 있는 생존환경 지표를 식별하는 데 목적이 있다.",
    )

    add_heading(doc, "4. 데이터와 시공간 모델링", 1)
    add_para(
        doc,
        "본 연구의 기본 분석 단위는 팀과 시점의 조합이다. PUBG의 최종 등수는 팀 단위로 결정되므로, 개별 플레이어보다 "
        "팀을 중심으로 분석하는 것이 적절하다. 한 시점에서 팀의 위치는 생존한 팀원들의 위치를 평균하여 계산한다. "
        "또한 팀원 간 거리를 이용하여 팀이 얼마나 흩어져 있었는지도 함께 고려한다.",
    )
    add_table(
        doc,
        ["자료", "본 연구에서의 의미"],
        [
            ("플레이어 위치", "각 시점의 플레이어 좌표와 생존 상태를 나타내며, 이동궤적의 기초가 된다."),
            ("팀 상태", "팀 중심 위치, 생존 인원, 팀원 간 분산을 포함한다."),
            ("안전구역 상태", "단계별 안전구역 중심, 반지름, 면적, 축소 비율을 포함한다."),
            ("전투 이벤트", "피해, 기절, 처치 등 특정 시간과 위치에서 발생한 사건을 나타낸다."),
            ("경기 결과", "팀 최종 등수와 생존 시간을 포함한다."),
        ],
        [2200, 7160],
    )
    add_para(
        doc,
        "안전구역은 중심과 반지름을 가진 원형 영역으로 다룬다. 각 시점에서 안전구역 면적을 계산하고, 이를 기준으로 "
        "전체 생존 팀 밀도와 팀별 위치 압박 정도를 분석한다. 이 구조를 통해 PUBG 텔레메트리는 이동 객체, 동적 공간 영역, "
        "공간 사건이 결합된 시공간 데이터베이스로 표현된다.",
    )

    add_heading(doc, "5. 연구 방법", 1)
    add_heading(doc, "5.1 데이터베이스 구성", 2)
    add_para(
        doc,
        "원시 텔레메트리는 경기 정보, 플레이어 위치, 팀 상태, 안전구역 상태, 전투 이벤트, 분석용 생존환경 자료로 나누어 "
        "정리한다. 경기 정보에는 맵 이름, 게임 모드, 경기 시간과 같은 기본 정보가 포함된다. 플레이어 위치 자료는 시간별 "
        "좌표와 체력, 생존 여부를 포함한다. 팀 상태 자료는 팀 중심 위치와 생존 인원, 팀원 간 분산 정도를 포함한다. "
        "안전구역 자료는 각 단계의 중심, 반지름, 면적을 포함하며, 전투 이벤트 자료는 발생 시각, 위치, 행위자와 대상자, "
        "피해량을 포함한다.",
    )
    add_heading(doc, "5.2 시공간 질의", 2)
    add_para(
        doc,
        "본 연구에서 사용하는 질의는 크게 다섯 가지이다. 안전구역 거리 질의는 각 팀이 안전구역 중심이나 경계로부터 "
        "얼마나 떨어져 있는지를 계산한다. 반경 질의는 특정 팀 주변 일정 거리 안에 있는 적 팀 수를 계산한다. 최근접 질의는 "
        "특정 팀과 가장 가까운 적 팀까지의 거리를 계산한다. 이동궤적 질의는 일정 시간 구간 동안 팀이 안전구역 안쪽으로 "
        "접근했는지 또는 경계 근처에 머물렀는지를 살펴본다. 시간창 질의는 특정 시점 이전 일정 시간 동안 팀이 관여한 "
        "전투 이벤트 수를 계산한다.",
    )
    add_table(
        doc,
        ["질의 유형", "분석 목적"],
        [
            ("안전구역 거리 질의", "팀이 안전구역 내부, 경계, 외부 중 어디에 있는지 판단한다."),
            ("반경 질의", "팀 주변 일정 거리 안에 있는 적 팀 수를 계산한다."),
            ("최근접 질의", "가장 가까운 적 팀까지의 거리를 계산한다."),
            ("이동궤적 질의", "안전구역 축소에 대한 팀의 이동 반응을 분석한다."),
            ("시간창 질의", "최근 교전 노출 정도를 계산한다."),
        ],
        [2500, 6860],
    )
    add_heading(doc, "5.3 생존환경 지표", 2)
    add_para(
        doc,
        "본 연구에서 생존환경은 한 팀이 특정 시점에 처한 공간적, 시간적 위험 상태를 의미한다. 주요 지표는 안전구역 경계와의 "
        "거리, 안전구역 외부 체류 여부, 안전구역 접근 속도, 주변 적 밀도, 가장 가까운 적과의 거리, 전체 팀 밀도, 최근 교전 "
        "횟수, 팀원 간 분산이다. 이 지표들은 단독으로 해석하기보다 함께 고려해야 한다. 예를 들어 안전구역 안에 있더라도 "
        "주변 적 밀도가 높으면 위험할 수 있고, 주변 적이 적더라도 안전구역 바깥에 오래 있으면 생존에 불리할 수 있다.",
    )

    add_heading(doc, "6. 분석 계획", 1)
    add_para(
        doc,
        "첫 번째 단계에서는 PUBG 텔레메트리를 시공간 데이터베이스 구조로 정리한다. 이 단계에서는 주요 객체 정의, "
        "분석 단위, 데이터베이스 구조, 질의 구조를 제시한다. 두 번째 단계에서는 안전구역 축소에 따른 생존환경 변화를 "
        "단계별로 분석한다. 단계가 진행될수록 안전구역 면적이 줄어들고, 생존 팀 밀도와 주변 적 밀도가 어떻게 변하는지 "
        "확인한다.",
    )
    add_para(
        doc,
        "세 번째 단계에서는 최종 등수와 생존환경 지표의 관계를 분석한다. 최종 등수는 숫자가 작을수록 좋은 성과이므로 "
        "상관 분석을 할 때 방향 해석에 주의한다. 추가로 상위 5위 또는 상위 10위 여부를 기준으로 분류 실험을 수행하여, "
        "어떤 지표가 상위권 진입과 관련이 큰지 확인한다. 마지막으로 기본적인 경기 요약 정보만 사용한 경우와 시공간 질의 "
        "기반 지표를 함께 사용한 경우를 비교한다.",
    )

    add_heading(doc, "7. 기대 결과와 기여", 1)
    add_para(
        doc,
        "본 연구를 통해 PUBG 텔레메트리가 시공간 데이터베이스로 표현될 수 있음을 보일 수 있을 것으로 예상한다. 특히 "
        "플레이어와 팀의 위치는 이동 객체로, 안전구역은 동적 공간 영역으로, 전투 이벤트는 시간과 위치를 가진 공간 사건으로 "
        "정리될 수 있다.",
    )
    add_para(
        doc,
        "안전구역 축소가 진행될수록 생존 팀이 더 좁은 공간에 모이면서 전체 팀 밀도는 증가하고, 가장 가까운 적과의 거리는 "
        "감소할 가능성이 있다. 이 과정에서 상위권 팀과 하위권 팀은 안전구역에 대한 반응에서 차이를 보일 수 있다. "
        "예를 들어 상위권 팀은 안전구역 외부 체류 시간이 짧거나, 경계 압박이 낮거나, 주변 적 밀도가 낮은 경로를 선택했을 "
        "가능성이 있다.",
    )
    add_para(
        doc,
        "연구의 기여는 세 가지로 정리된다. 첫째, PUBG 텔레메트리를 단순 로그 데이터가 아니라 시공간 데이터베이스로 "
        "정리한다. 둘째, 안전구역을 중심으로 팀의 이동궤적을 분석한다. 셋째, 최종 등수와 관련 있는 시공간 생존환경 지표를 "
        "식별한다.",
    )

    add_heading(doc, "8. 연구 한계", 1)
    add_para(
        doc,
        "첫째, 위치 데이터의 시간 해상도에 한계가 있을 수 있다. 위치가 연속적으로 기록되지 않고 일정 간격으로만 기록된다면 "
        "실제 이동 경로와 관측된 이동 경로 사이에 차이가 생길 수 있다. 둘째, 본 연구는 인과관계를 직접 증명하지 않는다. "
        "주변 적 밀도가 높고 최종 등수가 낮다고 해서 적 밀도가 낮은 등수를 직접 유발했다고 단정할 수는 없다.",
    )
    add_para(
        doc,
        "셋째, 분석 단위에 따라 해석이 달라질 수 있다. 솔로 모드에서는 플레이어 단위 분석이 적절하지만, 스쿼드 모드에서는 "
        "팀 단위 의사결정과 팀원 간 분산을 고려해야 한다. 넷째, 공식 텔레메트리만으로는 건물 내부 여부, 엄폐, 시야, 지형 "
        "조건을 정확히 알 수 없다. 따라서 본 연구는 맵 구조를 직접 해석하기보다, 공식 좌표와 안전구역, 전투 이벤트로 "
        "관측 가능한 시공간 환경만 분석한다.",
    )

    add_heading(doc, "9. 연구 일정", 1)
    add_table(
        doc,
        ["주차", "수행 내용"],
        [
            ("1주차", "PUBG 텔레메트리 구조 파악, 분석 단위 결정, 주요 테이블 설계"),
            ("2주차", "플레이어 위치, 팀 상태, 안전구역 상태, 전투 이벤트 전처리"),
            ("3주차", "안전구역 거리 질의, 반경 질의, 최근접 질의 구현"),
            ("4주차", "팀별 생존환경 지표 생성 및 단계별 변화 분석"),
            ("5주차", "최종 등수와 지표의 관계 분석, 상위권 분류 실험"),
            ("6주차", "결과 시각화, 한계 정리, 최종 보고서 작성"),
        ],
        [1500, 7860],
    )

    add_heading(doc, "참고문헌", 1)
    refs = [
        "Sagong, C., Woo, J. Y., & Kim, H. K. (2025). Are you a Hunter or Camper? A phase-based framework for real-time survival prediction in battle royale using FightScore and ActiveScore. Entertainment Computing, 55, 101040.",
        "PUBG Developer Documentation. Telemetry. https://documentation.pubg.com/en/telemetry.html",
        "PUBG Developer Documentation. Telemetry Events. https://documentation.pubg.com/en/telemetry-events.html",
        "PUBG Developer Documentation. Telemetry Objects. https://documentation.pubg.com/en/telemetry-objects.html",
        "Güting, R. H., Böhlen, M. H., Erwig, M., Jensen, C. S., Lorentzos, N. A., Schneider, M., & Vazirgiannis, M. (2000). A Foundation for Representing and Querying Moving Objects. ACM Transactions on Database Systems, 25(1), 1-42.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(ref)
        set_run_font(r, size=10)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
